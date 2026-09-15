"""Catálogo de frases de Excel. Los textos editables viven únicamente en el libro."""
import json
from pathlib import Path
from string import Formatter

from configuracion import ErrorConfiguracion, cargar_configuracion

ESQUEMA = Path(__file__).resolve().parent / 'recursos/textos.schema.json'
_catalogo = None


class Catalogo:
    def __init__(self, ruta, idioma, esquema=ESQUEMA):
        from openpyxl import load_workbook
        self.idioma = idioma
        self.frases = {}
        try:
            libro = load_workbook(ruta, read_only=True, data_only=False)
        except Exception as error:
            raise ErrorConfiguracion('textos.archivo_error', ruta=str(ruta), detalle=str(error)) from error
        try:
            nombre_hoja = next((nombre for nombre in libro.sheetnames
                                if nombre.casefold() == idioma.casefold()), None)
            if nombre_hoja is None:
                raise ErrorConfiguracion('textos.idioma_ausente', idioma=idioma)
            hoja = libro[nombre_hoja]
            filas = hoja.iter_rows(values_only=True)
            cabecera = next(filas, ())
            if not cabecera or cabecera[0] != 'clave' or 'texto' not in cabecera:
                raise ErrorConfiguracion('textos.formato_error', detalle='clave / texto')
            columnas = [str(c).strip() if c is not None else '' for c in cabecera]
            if len(set(columnas)) != len(columnas) or '' in columnas:
                raise ErrorConfiguracion('textos.columnas_error')
            for fila in filas:
                if all(v is None for v in fila):
                    continue
                clave = fila[0]
                if not isinstance(clave, str) or not clave.strip() or clave in self.frases:
                    raise ErrorConfiguracion('textos.formato_error', detalle=f'clave: {clave!r}')
                variantes = {}
                for columna, valor in zip(columnas[1:], fila[1:]):
                    if valor is not None:
                        if not isinstance(valor, str) or valor.startswith('='):
                            raise ErrorConfiguracion('textos.formato_error', detalle=f'{clave}.{columna}')
                        variantes[columna] = valor
                self.frases[clave] = variantes
        finally:
            libro.close()
        try:
            contrato = json.loads(Path(esquema).read_text(encoding='utf-8'))
        except (OSError, ValueError) as error:
            raise ErrorConfiguracion('textos.formato_error', detalle=str(error)) from error
        normalizadas = {}
        for clave, variantes in contrato.items():
            normalizadas[clave] = {}
            for indice, (variante, parametros) in enumerate(variantes.items()):
                columna = 'texto' if indice == 0 else f'variante_{indice}'
                valor = self.frases.get(clave, {}).get(columna)
                if valor is None:
                    raise ErrorConfiguracion('textos.frase_ausente', clave=f'{clave}.{variante}', idioma=idioma)
                try:
                    partes = list(Formatter().parse(valor))
                    campos = {campo for _, campo, _, _ in partes if campo is not None}
                    if any('{' in formato or '}' in formato for _, _, formato, _ in partes if formato):
                        raise ValueError('nested_format')
                except ValueError as error:
                    raise ErrorConfiguracion('textos.formato_error', detalle=f'{clave}.{variante}') from error
                if campos != set(parametros):
                    raise ErrorConfiguracion('textos.parametros_error', clave=f'{clave}.{variante}', parametros=', '.join(parametros))
                normalizadas[clave][variante] = valor
        self.frases = normalizadas

    def texto(self, identificador, variante='texto', **datos):
        clave = identificador
        variantes = self.frases.get(clave, {})
        if variante == 'texto' and 'cantidad' in datos and 'singular' in variantes and 'plural' in variantes:
            variante = 'singular' if datos['cantidad'] == 1 else 'plural'
        try:
            plantilla = variantes[variante]
        except KeyError as error:
            raise ErrorConfiguracion('textos.frase_ausente', clave=f'{clave}.{variante}', idioma=self.idioma) from error
        try:
            return plantilla.format(**datos)
        except (KeyError, ValueError, IndexError, AttributeError) as error:
            raise ErrorConfiguracion('textos.formato_error', detalle=f'{clave}.{variante}: {error}') from error


def iniciar_textos(configuracion=None):
    global _catalogo
    config = configuracion or cargar_configuracion()
    nuevo = Catalogo(config.textos, config.languaje)
    _catalogo = nuevo
    return nuevo


def tr(identificador, variante='texto', **datos):
    if _catalogo is None:
        iniciar_textos()
    return _catalogo.texto(identificador, variante, **datos)
