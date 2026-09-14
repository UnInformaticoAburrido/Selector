"""Importación de nombres desde la primera columna de Excel o CSV."""
import csv
from pathlib import Path
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError


def importar_nombres(ruta, omitir_cabecera=False):
    """Lee la columna A de la primera pestaña (Excel) o columna inicial (CSV).

    Conserva el orden y los nombres repetidos. Omite celdas vacías y,
    únicamente si se solicita, la primera fila física del archivo.
    """
    ruta = Path(ruta)
    extension = ruta.suffix.lower()
    if extension == ".xlsx":
        valores = _leer_excel(ruta)
    elif extension == ".csv":
        valores = _leer_csv(ruta)
    else:
        raise ValueError("Formato no compatible. Selecciona un archivo Excel (.xlsx) o CSV (.csv).")
    if omitir_cabecera:
        valores = valores[1:]
    nombres = [str(valor).strip() for valor in valores if valor is not None and str(valor).strip()]
    if not nombres:
        raise ValueError("No se encontraron nombres en la primera columna.")
    return nombres


def _leer_excel(ruta):
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as error:
        raise ValueError(
            "Para importar Excel instala las dependencias:\n"
            "Windows: py -m pip install -r requirements.txt\n"
            "Linux: python3 -m pip install -r requirements.txt\n"
            "Usa el mismo entorno de Python con el que abres la aplicación."
        ) from error
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
        try:
            if not libro.worksheets:
                raise ValueError("El archivo no contiene ninguna hoja de cálculo.")
            hoja = libro.worksheets[0]
            # No depender de dimensiones incorrectas exportadas por otras aplicaciones.
            hoja.reset_dimensions()
            return [fila[0] for fila in hoja.iter_rows(min_col=1, max_col=1, values_only=True)]
        finally:
            libro.close()
    except (BadZipFile, InvalidFileException, ParseError, KeyError) as error:
        raise ValueError("No se pudo leer el Excel. Comprueba que sea un archivo .xlsx válido y sin contraseña.") from error


def _leer_csv(ruta):
    with ruta.open(encoding="utf-8-sig", newline="") as archivo:
        muestra = archivo.read(8192)
        archivo.seek(0)
        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t")
        except csv.Error:
            dialecto = csv.excel
        return [fila[0] if fila else None for fila in csv.reader(archivo, dialecto)]
