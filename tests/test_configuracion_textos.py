import ast
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from apariencia import detectar_tema_sistema
from configuracion import RUTA_CONFIGURACION, cargar_configuracion, ErrorConfiguracion
from textos import Catalogo, ESQUEMA

RAIZ = Path(__file__).resolve().parents[1]


class ConfiguracionTextosTests(unittest.TestCase):
    def setUp(self):
        self.temporal = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporal.cleanup)
        self.base = Path(self.temporal.name)
        self.ruta = self.base / 'configuracion.json'
        self.libro = self.base / 'textos.xlsx'
        shutil.copyfile(RAIZ / 'recursos/textos.xlsx', self.libro)
        self.datos = json.loads(RUTA_CONFIGURACION.read_text())
        self.datos['textos'] = 'textos.xlsx'
        self.escribir()

    def escribir(self):
        self.ruta.write_text(json.dumps(self.datos), encoding='utf-8')

    def modificar(self, clave, columna, valor):
        wb = load_workbook(self.libro)
        ws = wb['ES']
        for fila in ws.iter_rows(min_row=2):
            if fila[0].value == clave:
                ws.cell(fila[0].row, columna, valor)
                break
        wb.save(self.libro)
        wb.close()

    def test_config_resuelve_rutas_y_valores(self):
        self.datos.update(languaje='es', tema='Oscuro')
        self.datos['ventana'] = {'ancho':1000, 'alto':650, 'modo':'pantalla completa sin bordes'}
        self.escribir()
        config = cargar_configuracion(self.ruta)
        self.assertEqual(config.textos, self.libro.resolve())
        self.assertEqual((config.languaje,config.tema,config.ancho,config.alto,config.modo), ('ES','oscuro',1000,650,'sin_bordes'))

    def test_valores_invalidos_tienen_diagnostico(self):
        for campo, valor in [('tema','sepia'), ('languaje',None), ('ventana',None)]:
            original = self.datos[campo]
            self.datos[campo] = valor
            self.escribir()
            with self.subTest(campo=campo), self.assertRaises(ErrorConfiguracion):
                cargar_configuracion(self.ruta)
            self.datos[campo] = original
        for valor in [True, '1280', 0, 499, 20000]:
            self.datos['ventana']['ancho'] = valor
            self.escribir()
            with self.assertRaises(ErrorConfiguracion):
                cargar_configuracion(self.ruta)

    def test_json_corrupto(self):
        self.ruta.write_text('{')
        with self.assertRaises(ErrorConfiguracion) as captura:
            cargar_configuracion(self.ruta)
        self.assertEqual(captura.exception.clave, 'config.archivo_error')

    def test_variantes_y_plurales(self):
        catalogo = Catalogo(self.libro, 'ES')
        self.assertEqual(catalogo.texto('participante.cantidad', cantidad=1), '1 persona')
        self.assertEqual(catalogo.texto('participante.cantidad', cantidad=2), '2 personas')
        self.assertIn('ya había participado', catalogo.texto('seleccion.detalle','alternativa',numero=3,solicitado=1))

    def test_cambiar_pestana_y_editar_texto(self):
        wb=load_workbook(self.libro)
        hoja=wb.copy_worksheet(wb['ES']);hoja.title='EN'
        for fila in hoja.iter_rows(min_row=2):
            if fila[0].value == 'seleccion.accion':
                fila[1].value='Select'
        wb.save(self.libro);wb.close()
        self.datos['languaje']='EN';self.escribir()
        config=cargar_configuracion(self.ruta)
        self.assertEqual(Catalogo(config.textos,config.languaje).texto('seleccion.accion'), 'Select')
        self.assertEqual(Catalogo(self.libro,'ES').texto('seleccion.accion'), 'Seleccionar')

    def test_idioma_ausente_no_se_sustituye_silenciosamente(self):
        with self.assertRaises(ErrorConfiguracion) as captura:
            Catalogo(self.libro, 'FR')
        self.assertEqual(captura.exception.clave, 'textos.idioma_ausente')

    def test_marcador_incorrecto(self):
        self.modificar('participante.cantidad', 2, '{otro} personas')
        with self.assertRaises(ErrorConfiguracion) as captura:
            Catalogo(self.libro,'ES')
        self.assertEqual(captura.exception.clave, 'textos.parametros_error')

    def test_clave_borrada_y_duplicada(self):
        wb=load_workbook(self.libro);ws=wb['ES'];ws.delete_rows(2);wb.save(self.libro);wb.close()
        with self.assertRaises(ErrorConfiguracion) as captura:
            Catalogo(self.libro,'ES')
        self.assertEqual(captura.exception.clave, 'textos.frase_ausente')
        shutil.copyfile(RAIZ/'recursos/textos.xlsx',self.libro)
        wb=load_workbook(self.libro);ws=wb['ES'];ws.append([c.value for c in ws[2]]);wb.save(self.libro);wb.close()
        with self.assertRaises(ErrorConfiguracion):
            Catalogo(self.libro,'ES')

    def test_formula_no_es_texto(self):
        self.modificar('app.nombre', 2, '=1+1')
        with self.assertRaises(ErrorConfiguracion):
            Catalogo(self.libro,'ES')

    def test_todas_las_claves_del_codigo_existen_en_excel(self):
        catalogo=Catalogo(self.libro,'ES')
        esquema=json.loads(ESQUEMA.read_text())
        for ruta in RAIZ.glob('*.py'):
            for nodo in ast.walk(ast.parse(ruta.read_text())):
                if isinstance(nodo,ast.Call) and isinstance(nodo.func,ast.Name) and nodo.func.id=='tr':
                    if not nodo.args or not isinstance(nodo.args[0],ast.Constant):
                        continue
                    clave=nodo.args[0].value
                    self.assertIn(clave,catalogo.frases,(ruta.name,nodo.lineno))
                    if len(nodo.args)>1 and isinstance(nodo.args[1],ast.Constant):
                        self.assertIn(nodo.args[1].value,esquema[clave])

    def test_sistema_detecta_claro_oscuro_y_fallback(self):
        for valor, esperado in [('1','oscuro'),('2','claro')]:
            with patch('apariencia.sys.platform','linux'), patch('apariencia.subprocess.run',return_value=subprocess.CompletedProcess([],0,f'(<uint32 {valor}>,)','')):
                self.assertEqual(detectar_tema_sistema(),esperado)
        with patch('apariencia.sys.platform','linux'), patch('apariencia.subprocess.run',side_effect=FileNotFoundError):
            self.assertEqual(detectar_tema_sistema(),'claro')


if __name__ == '__main__':
    unittest.main()
